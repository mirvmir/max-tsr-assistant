from hashlib import sha256
from io import BytesIO
from uuid import uuid4
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from tsr.adapters.render import render_document
from tsr.contracts import DocumentField, DocumentModel, DocumentSection, RenderRequest, VersionRef


def demo_model():
    return DocumentModel(
        schema_version='1.0.0', document_kind='application', case_mode='demo',
        manifest_hash='a' * 64, missing_fields=('applicant_name',), sources=(),
        warnings=('Модельный маршрут не является утверждённым обращением.',),
        sections=(DocumentSection(title='Черновик обращения', text='Не является официальной формой.',
                  fields=(DocumentField(label='Заявитель', value='Не указано'),
                          DocumentField(label='Пояснение', value='<script>Кириллица</script> ' * 100))),),
    )


def request(model, fmt='pdf', limit=2_000_000):
    return RenderRequest(job_id=uuid4(), fence_token=2, manifest_hash=model.manifest_hash,
                         document_kind=model.document_kind, format=fmt,
                         template_ref=VersionRef(id='demo-application', version='1.0.0'),
                         model=model, max_output_bytes=limit)


def test_render_cyrillic_literal_markup_docx_and_byte_bound():
    model = demo_model()
    result = render_document(request(model, 'docx'))
    assert result.ok, result.error
    rendered = result.value
    assert rendered.plaintext_sha256 == sha256(rendered.bytes).hexdigest()
    paragraphs = '\n'.join(p.text for p in Document(BytesIO(rendered.bytes)).paragraphs)
    assert '<script>Кириллица</script>' in paragraphs
    assert 'Не указано' in paragraphs
    assert 'ДЕМОНСТРАЦИОННЫЙ' in paragraphs
    assert 'ФИО заявителя: Не указано' in paragraphs
    assert document_title_has_no_border(Document(BytesIO(rendered.bytes)))
    assert not render_document(request(model, 'pdf', limit=100)).ok


def document_title_has_no_border(document):
    from docx.oxml.ns import qn
    return not list(document.styles['Title']._element.iter(qn('w:pBdr')))


def test_pdf_has_cyrillic_font_and_binds_manifest_hash():
    model = demo_model()
    result = render_document(request(model))
    assert result.ok, result.error
    assert result.value.bytes.startswith(b'%PDF-')
    assert b'DejaVuSans' in result.value.bytes
    mismatched = request(model).model_copy(update={'manifest_hash': 'b' * 64})
    assert not render_document(mismatched).ok
    pilot = demo_model().model_copy(update={'case_mode':'pilot'})
    assert not render_document(request(pilot)).ok


def test_invalid_table_validation_is_a_safe_typed_report():
    from tsr.domain.documents import validate_document_model
    from tsr.contracts import DocumentTable
    invalid = demo_model().model_copy(update={'sections':(DocumentSection(
        tables=(DocumentTable(headers=('A','B'),rows=(('single',),)),)),)})
    report = validate_document_model(invalid, None)
    assert not report.valid
    assert report.errors[0].code == 'invalid_table_shape'


def frozen_manifest(branch='support'):
    from tsr.contracts import (CalculatedEvidence, ComparisonResult, DocumentManifest, Fact,
        InputRevision, ManifestContent, OfferSnapshot, PricingResult, RouteEvaluation, SourceMetadata,
        content_hash)
    import json
    root = Path(__file__).resolve().parents[1]
    offer = OfferSnapshot.model_validate(json.loads((root / 'data/catalog/mvp/1.0.0/catalog.json').read_text())['offers'][0])
    now = datetime(2026, 9, 29, tzinfo=timezone.utc)
    owner, case, revision = uuid4(), uuid4(), uuid4()
    inp = InputRevision(input_revision_id=revision, owner_id=owner, case_id=case, created_at=now,
                        category_id=offer.variant.category_id, profile_ref=offer.profile_ref,
                        role='representative', region_code=None, document_fields={'address':'Длинный адрес ' * 30})
    matching, pricing_ref = VersionRef(id='matching-v1',version='1.0.0'),VersionRef(id='pricing-v1',version='1.0.0')
    comparison=ComparisonResult(comparison_id=uuid4(),input_revision_id=revision,snapshot_id=offer.snapshot_id,
        profile_ref=offer.profile_ref,matching_algorithm_ref=matching,computed_at=now,fields=(),classification='incomplete')
    pricing=PricingResult(pricing_algorithm_ref=pricing_ref,input_revision_id=revision,snapshot_id=offer.snapshot_id,
        certificate_use='unknown',delivery=offer.delivery,status='incomplete',
        provenance=CalculatedEvidence(algorithm=pricing_ref,input_refs=(str(revision),)))
    route=RouteEvaluation(status='needs_clarification',addressee=Fact.unknown(),evaluated_at=now,
                         missing_fields=('region_code',)) if branch=='support' else None
    template_ids=['demo-purchase-card'] if branch=='purchase' else ['demo-application','demo-product-card','demo-checklist']
    content=ManifestContent(owner_id=owner,case_id=case,case_revision=2,deletion_epoch=0,case_mode='demo',
        input_revision=inp,offer_snapshot=offer,sources=(SourceMetadata(source_id='synthetic',title='Синтетический источник'),),
        comparison=comparison,pricing=pricing,route=route,branch=branch,category_profile_ref=offer.profile_ref,
        matching_algorithm_ref=matching,pricing_algorithm_ref=pricing_ref,
        template_refs=tuple(VersionRef(id=key,version='1.0.0') for key in template_ids),
        generator_ref=VersionRef(id='generator-v1',version='1.0.0'),release_commit='frozen-release',missing_fields=('region_code',))
    return DocumentManifest(manifest_id=uuid4(),confirmation_id=uuid4(),created_at=now,
                            manifest_hash=content_hash(content),content=content)


def test_bundle_branch_composition_frozen_content_and_unknowns():
    from tsr.domain.documents import build_document_model, build_document_request
    purchase = frozen_manifest('purchase')
    assert [(x.document_kind,x.format) for x in build_document_request(purchase).required_artifacts] == [('purchase_card','pdf')]
    support = frozen_manifest()
    assert [(x.document_kind,x.format) for x in build_document_request(support).required_artifacts] == [
        ('application','docx'),('application','pdf'),('product_card','pdf'),('checklist','pdf')]
    model = build_document_model(support,'application')
    assert model.ok, model.error
    assert 'applicant_name' in model.value.missing_fields
    assert 'addressee' in model.value.missing_fields
    assert 'region_code' in model.value.missing_fields
    text = '\n'.join(p.value for s in model.value.sections for p in s.fields)
    assert 'Представитель' in text
    assert 'Неизвестно' in text
    assert 'frozen-release' in text
    assert any('Модельный маршрут' in warning for warning in model.value.warnings)
    changed = support.model_copy(update={'manifest_hash':'b'*64})
    assert not build_document_model(changed,'application').ok


def test_public_offer_facts_are_not_described_as_synthetic_prices():
    from tsr.contracts import content_hash, Supplier
    from tsr.domain.documents import build_document_model
    frozen = frozen_manifest('purchase')
    public_offer = frozen.content.offer_snapshot.model_copy(update={'data_kind':'public_snapshot'})
    supplier=Supplier(supplier_id=public_offer.supplier_id,display_name='Поставщик публичного предложения')
    content = frozen.content.model_copy(update={'offer_snapshot':public_offer,'supplier':supplier})
    manifest = frozen.model_copy(update={'content':content,'manifest_hash':content_hash(content)})
    model = build_document_model(manifest,'purchase_card').value
    warnings = '\n'.join(model.warnings)
    assert 'публичного снимка' in warnings
    assert 'Все предложения и условия являются синтетическими' not in warnings
    assert any(p.value == '120 000,00 ₽' for s in model.sections for p in s.fields)
    assert any(p.value == supplier.display_name for s in model.sections for p in s.fields)


def test_seller_code_and_conditional_money_are_explicit_in_rendered_text():
    from tsr.contracts import Money, content_hash
    from tsr.domain.documents import build_document_model
    from tsr.domain.pricing import quote_purchase
    frozen = frozen_manifest('purchase')
    inp = frozen.content.input_revision.model_copy(update={'certificate_amount': Money(minor=10_000_000)})
    pricing = quote_purchase(inp, frozen.content.offer_snapshot, frozen.content.pricing_algorithm_ref).value
    assert pricing.certificate_use == 'conditional'
    content = frozen.content.model_copy(update={'input_revision': inp, 'pricing': pricing})
    manifest = frozen.model_copy(update={'content': content, 'manifest_hash': content_hash(content)})
    model = build_document_model(manifest, 'purchase_card').value
    fields = {p.label: p.value for s in model.sections for p in s.fields}
    assert fields['Код предложения продавца'] == content.offer_snapshot.seller_sku
    assert 'Артикул' not in fields
    assert fields['Условное покрытие сертификатом'] == '100 000,00 ₽'
    assert fields['Условная разница без доставки'] == '20 000,00 ₽'
    assert 'Условные собственные расходы' in fields
    from tsr.adapters.render import _plain_lines
    # These lines are the shared projection consumed by both renderers.
    text = '\n'.join(_plain_lines(model))
    assert 'Код предложения продавца: ' + content.offer_snapshot.seller_sku in text
    assert 'Условное покрытие сертификатом: 100 000,00 ₽' in text
    assert 'Условная разница без доставки: 20 000,00 ₽' in text
    render_request = request(model).model_copy(update={'template_ref': VersionRef(id='demo-purchase-card', version='1.0.0')})
    assert render_document(render_request).ok


def renamed_registry():
    import json
    from tsr.contracts import TemplateRegistry
    root = Path(__file__).resolve().parents[1]
    registry = TemplateRegistry.model_validate(json.loads((root/'templates/registry.json').read_text()))
    return registry.model_copy(update={'templates': tuple(entry.model_copy(update={
        'template_id': 'next-' + entry.document_kind, 'version': '2.0.0'}) for entry in registry.templates)})


def test_registry_binds_new_template_versions_and_rejects_unbound_or_unknown_renderer():
    from tsr.contracts import content_hash
    from tsr.domain.documents import build_document_request, build_document_model
    registry = renamed_registry()
    frozen = frozen_manifest('purchase')
    entry = registry.templates[0]
    content = frozen.content.model_copy(update={'template_refs': (entry.ref,)})
    manifest = frozen.model_copy(update={'content': content, 'manifest_hash': content_hash(content)})
    spec = build_document_request(manifest, registry=registry).required_artifacts[0]
    assert spec.template_ref == entry.ref
    import pytest
    with pytest.raises(ValueError, match='Frozen template reference'):
        build_document_request(frozen, registry=registry)
    model = build_document_model(manifest, 'purchase_card').value
    bound = request(model).model_copy(update={'template_ref': entry.ref, 'template_registry': registry})
    assert render_document(bound).ok
    assert not render_document(bound.model_copy(update={'template_ref': VersionRef(id=entry.template_id, version='2.0.1')})).ok
    bad_entry = entry.model_copy(update={'assets': (entry.assets[0].model_copy(update={'renderer_id': 'arbitrary-program'}),)})
    bad_registry = registry.model_copy(update={'templates': (bad_entry,)})
    assert not render_document(bound.model_copy(update={'template_registry': bad_registry})).ok
    wrong_kind = registry.model_copy(update={'templates': (entry.model_copy(update={'document_kind': 'application'}),)})
    assert not render_document(bound.model_copy(update={'template_registry': wrong_kind})).ok


def test_pilot_renderer_requires_fresh_public_reviewed_registry():
    from datetime import timedelta
    from tsr.contracts import Review
    now = datetime.now(timezone.utc)
    registry = renamed_registry()
    entry = next(e for e in registry.templates if e.document_kind == 'application')
    model = demo_model().model_copy(update={'case_mode': 'pilot'})
    bound = request(model, 'docx').model_copy(update={'template_ref': entry.ref, 'template_registry': registry})
    assert not render_document(bound).ok
    reviewed = registry.model_copy(update={'data_kind': 'public_snapshot', 'review': Review(status='reviewed',
        reviewer_id='operator', reviewed_at=now-timedelta(days=1), review_due_at=now+timedelta(days=1))})
    result = render_document(bound.model_copy(update={'template_registry': reviewed}))
    assert result.ok, result.error
    text = '\n'.join(p.text for p in Document(BytesIO(result.value.bytes)).paragraphs)
    assert 'ДЕМОНСТРАЦИОННЫЙ' not in text
    assert render_document(bound.model_copy(update={'template_registry': reviewed, 'format': 'pdf'})).ok
    expired = reviewed.model_copy(update={'review': Review(status='reviewed', reviewer_id='operator',
        reviewed_at=now-timedelta(days=2), review_due_at=now-timedelta(days=1))})
    assert not render_document(bound.model_copy(update={'template_registry': expired})).ok
    assert not render_document(bound.model_copy(update={'template_registry': reviewed.model_copy(update={'review': Review()})})).ok


def test_pricing_explanations_and_supplier_questions_are_human_readable():
    from tsr.contracts import Delivery, Fact, Money, content_hash
    from tsr.domain.documents import build_document_model
    from tsr.domain.pricing import quote_purchase
    frozen = frozen_manifest('purchase')
    inp = frozen.content.input_revision.model_copy(update={'certificate_amount': Money(minor=10_000_000)})
    offer = frozen.content.offer_snapshot.model_copy(update={'delivery': Delivery(
        mode='unknown', charge=Fact.unknown(), terms=Fact.unknown())})
    pricing = quote_purchase(inp, offer, frozen.content.pricing_algorithm_ref).value
    content = frozen.content.model_copy(update={'input_revision': inp, 'offer_snapshot': offer, 'pricing': pricing})
    manifest = frozen.model_copy(update={'content': content, 'manifest_hash': content_hash(content)})
    model = build_document_model(manifest, 'purchase_card').value
    from tsr.adapters.render import _plain_lines
    text = '\n'.join(_plain_lines(model))
    assert 'Условия и стоимость доставки неизвестны' in text
    assert 'Условия применения сертификата не подтверждены' in text
    assert 'Расчёт предполагает, что сертификат применим и продавец его принимает' in text
    assert 'Уточните условия и стоимость доставки' in text
    assert 'Уточните, принимает ли продавец сертификат' in text
    for code in (*pricing.reasons, *pricing.assumptions):
        assert code not in text


def test_route_actions_and_known_missing_fields_are_human_readable():
    from tsr.contracts import content_hash
    from tsr.domain.documents import build_document_model
    from tsr.adapters.render import _plain_lines
    frozen = frozen_manifest()
    steps = ('route.synthetic_model', 'route.verify_real_rules', 'route.no_approval_promise')
    missing = ('supplier_accepts_certificate', 'supplier_has_fund_contract', 'applicant_status_declared', 'certificate_applicable_declared')
    route = frozen.content.route.model_copy(update={'next_steps': steps, 'missing_fields': missing})
    content = frozen.content.model_copy(update={'route': route})
    manifest = frozen.model_copy(update={'content': content, 'manifest_hash': content_hash(content)})
    text = '\n'.join(_plain_lines(build_document_model(manifest, 'checklist').value))
    for code in (*steps, *missing):
        assert code not in text
    assert 'Проверьте актуальные требования у адресата' in text
    assert 'Одобрение обращения не гарантируется' in text
    assert 'Продавец принимает сертификат: Не указано' in text
    assert 'Договор продавца с фондом: Не указано' in text
