"""In-memory renderer entrypoint safe to import into an explicit spawn child.

Only content DTOs, immutable local assets and formatting libraries are loaded.
No database, configuration, transport or encryption dependency is imported.
"""
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from tsr.contracts import RenderedBytes, RenderRequest, Result


ASSET_ROOT = Path(__file__).resolve().parents[4]
DEMO_MARK = 'ДЕМОНСТРАЦИОННЫЙ ЧЕРНОВИК'
MISSING_LABELS = {
    'applicant_name':'ФИО заявителя', 'representative_name':'ФИО представителя',
    'region_code':'Регион', 'addressee':'Адресат', 'route':'Маршрут обращения',
    'address':'Адрес', 'phone':'Телефон', 'email':'Электронная почта',
    'applicant_status_declared':'Статус заявителя',
    'certificate_applicable_declared':'Применимость сертификата',
}
SUPPORTED = {
    ('purchase_card', 'pdf'): 'demo-purchase-card',
    ('application', 'pdf'): 'demo-application',
    ('application', 'docx'): 'demo-application',
    ('product_card', 'pdf'): 'demo-product-card',
    ('checklist', 'pdf'): 'demo-checklist',
}


def _plain_lines(model):
    if model.case_mode == 'demo':
        yield DEMO_MARK
        yield 'Синтетические данные. Документ не является официальной формой.'
    for warning in model.warnings:
        yield warning
    for section in model.sections:
        if section.title:
            yield section.title
        if section.text:
            yield section.text
        for field in section.fields:
            yield f'{field.label}: {field.value}'
        for table in section.tables:
            yield ' | '.join(table.headers)
            for row in table.rows:
                yield ' | '.join(row)
    if model.missing_fields:
        yield 'Незаполненные поля'
        for field in model.missing_fields:
            yield f'{MISSING_LABELS.get(field, field)}: Не указано'
    if model.sources:
        yield 'Источники'
        for source in model.sources:
            yield f'{source.title} ({source.source_id})'
            if source.url:
                yield source.url
            if source.locator:
                yield source.locator
    yield f'Хеш подтверждённого содержания: {model.manifest_hash}'


def _pdf(model):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    pdfmetrics.registerFont(TTFont('DejaVuSans', str(ASSET_ROOT / 'assets/fonts/DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', str(ASSET_ROOT / 'assets/fonts/DejaVuSans-Bold.ttf')))
    normal = ParagraphStyle('body', fontName='DejaVuSans', fontSize=10, leading=15,
                            spaceAfter=7, splitLongWords=True)
    heading = ParagraphStyle('heading', parent=normal, fontName='DejaVuSans-Bold',
                             fontSize=13, leading=18, spaceBefore=8, keepWithNext=True)
    headings = {s.title for s in model.sections} | {DEMO_MARK, 'Незаполненные поля', 'Источники'}
    story = []
    for line in _plain_lines(model):
        # Literal input never becomes reportlab markup or a filesystem reference.
        story.append(Paragraph(escape(line).replace('\n', '<br/>'), heading if line in headings else normal))
    output = BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('DejaVuSans', 8)
        canvas.setFillColor(colors.HexColor('#555555'))
        if model.case_mode == 'demo':
            canvas.drawString(42, 26, DEMO_MARK)
        canvas.drawRightString(A4[0] - 42, 26, str(doc.page))
        canvas.restoreState()

    SimpleDocTemplate(output, pagesize=A4, rightMargin=42, leftMargin=42,
                      topMargin=40, bottomMargin=48, title='Материалы по ТСР',
                      author='', pageCompression=1).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _docx(model):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    document.core_properties.author = ''
    document.core_properties.title = 'Черновик обращения по ТСР'
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    for name in ('Normal', 'Title', 'Heading 1'):
        style = document.styles[name]
        style.font.name = 'DejaVu Sans'
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    document.styles['Normal'].font.size = Pt(10)
    document.styles['Title'].font.size = Pt(16)
    document.styles['Heading 1'].font.size = Pt(12)
    document.styles['Normal'].paragraph_format.space_after = Pt(7)
    headings = {s.title for s in model.sections} | {'Незаполненные поля', 'Источники'}
    for index, line in enumerate(_plain_lines(model)):
        style = 'Title' if index == 0 else 'Heading 1' if line in headings else 'Normal'
        document.add_paragraph(line, style)
    if model.case_mode == 'demo':
        section.footer.paragraphs[0].text = DEMO_MARK
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_document(request: RenderRequest) -> Result[RenderedBytes]:
    model = request.model
    expected_template = SUPPORTED.get((request.document_kind, request.format))
    if (model.manifest_hash != request.manifest_hash or model.document_kind != request.document_kind
            or expected_template is None or request.template_ref.id != expected_template
            or request.template_ref.version != '1.0.0' or request.max_output_bytes <= 0):
        return Result.failure('VALIDATION_ERROR', safe_message_key='document_invalid')
    if model.schema_version != '1.0.0':
        return Result.failure('UNSUPPORTED_SCHEMA')
    # Every current registered asset is synthetic and draft. A pilot must wait
    # for separately reviewed templates, never hide the demo mark on these.
    if model.case_mode != 'demo':
        return Result.failure('DATA_NOT_READY', safe_message_key='template_not_reviewed')
    # Bound input before formatting to avoid pathological unbounded in-memory work.
    if sum(len(line.encode('utf-8')) for line in _plain_lines(model)) > 1_000_000:
        return Result.failure('VALIDATION_ERROR', safe_message_key='document_too_large')
    try:
        data = _pdf(model) if request.format == 'pdf' else _docx(model)
        if len(data) > min(request.max_output_bytes, 10_485_760):
            return Result.failure('VALIDATION_ERROR', safe_message_key='document_too_large')
        return Result.success(RenderedBytes(
            job_id=request.job_id, fence_token=request.fence_token,
            manifest_hash=request.manifest_hash, format=request.format,
            mime_type='application/pdf' if request.format == 'pdf' else
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            plaintext_sha256=sha256(data).hexdigest(), bytes=data,
        ))
    except Exception:
        # Do not leak render-library messages containing submitted content or paths.
        return Result.failure('PERMANENT_FAILURE', safe_message_key='document_render_failed')
